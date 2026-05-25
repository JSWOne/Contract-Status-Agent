"""
Power Automate-backed SKU master lookup helper.

The Cloud Run service does not read the SharePoint Excel master directly. It
calls Power Automate helper flows that own Excel/SharePoint access and return
filtered material/SKU data.
"""

from __future__ import annotations

import os
import json
import ast
import re
from pathlib import Path
from typing import Any

import requests
from dotenv import load_dotenv


load_dotenv(Path(__file__).with_name(".env"))


def normalise_party_code(value: str) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    if not digits:
        return ""
    return digits.zfill(10)


def normalise_plant_code(value: str) -> str:
    return re.sub(r"\D", "", str(value or ""))


def get_sku_choices(division: str, bp_code: str, sp_code: str, ship_plant_code: str = "") -> dict[str, Any]:
    plant_code = normalise_plant_code(ship_plant_code)
    payload = {
        "action": "get_sku_choices",
        "division": division,
        "bp_code": normalise_party_code(bp_code),
        "sp_code": normalise_party_code(sp_code),
        "ship_plant_code": plant_code,
        "ship_plant": plant_code,
        "plant_code": plant_code,
    }
    data = _call_lookup(payload)
    if data.get("skus") or data.get("rows"):
        return data
    if data.get("materials"):
        return _with_second_stage_skus(data, payload)

    # GI helper flow was built with `get_materials` as the first-stage action.
    # Retry with that action when no choices are returned.
    payload["action"] = "get_materials"
    data = _call_lookup(payload)
    if data.get("skus") or data.get("rows"):
        return data
    return _with_second_stage_skus(data, payload)


def _with_second_stage_skus(data: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
    # Some division flows (e.g. GI) expose SKUs via a second-stage get_skus action.
    materials = _extract_material_values(data.get("materials"))
    if not materials:
        return data

    all_skus: list[Any] = []
    for material in materials:
        sku_payload = dict(payload)
        sku_payload["action"] = "get_skus"
        sku_payload["material"] = material
        sku_data = _call_lookup(sku_payload)
        all_skus.extend(sku_data.get("skus") or [])

    deduped = _dedupe_items(all_skus)
    if deduped:
        data["skus"] = deduped
    return data


def _extract_material_values(materials: Any) -> list[str]:
    values: list[str] = []
    for item in materials or []:
        if isinstance(item, dict):
            value = str(item.get("material") or item.get("value") or "").strip()
        else:
            value = str(item or "").strip()
        if value:
            values.append(value)
    return list(dict.fromkeys(values))


def _dedupe_items(items: list[Any]) -> list[Any]:
    seen: set[str] = set()
    output: list[Any] = []
    for item in items:
        key = repr(item)
        if key in seen:
            continue
        seen.add(key)
        output.append(item)
    return output


def get_sku_details(
    division: str,
    bp_code: str,
    sp_code: str,
    ship_plant_code: str,
    material: str,
    description: str,
) -> dict[str, Any]:
    plant_code = normalise_plant_code(ship_plant_code)
    payload = {
        "action": "get_sku_details",
        "division": division,
        "bp_code": normalise_party_code(bp_code),
        "sp_code": normalise_party_code(sp_code),
        "ship_plant_code": plant_code,
        "ship_plant": plant_code,
        "plant_code": plant_code,
        "material": (material or "").strip(),
        "description": (description or "").strip(),
    }
    data = _call_lookup(payload)

    # GI details lookup in PA can intermittently fail with upstream 502 and return
    # empty rows. When that happens, use an optional local Excel fallback so local
    # testing can proceed without affecting HRC/CRCA flows.
    if (
        str(division or "").strip().upper() == "GI"
        and not (data.get("rows") or [])
    ):
        row = _gi_details_from_local_master(payload)
        if row:
            data["rows"] = [row]
    return data


def _gi_details_from_local_master(payload: dict[str, Any]) -> dict[str, Any]:
    master_path = os.getenv("GI_MASTER_XLSX_PATH", "").strip()
    if not master_path:
        return {}
    path = Path(master_path)
    if not path.exists():
        return {}
    try:
        import openpyxl
    except Exception:
        return {}

    target_bp = normalise_party_code(payload.get("bp_code", ""))
    target_sp = normalise_party_code(payload.get("sp_code", ""))
    target_plant = normalise_plant_code(payload.get("plant_code", ""))
    target_material = str(payload.get("material") or "").strip().upper()
    target_desc = str(payload.get("description") or "").strip().upper()

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        header = [str(c or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        idx = {name: i for i, name in enumerate(header)}
        required = ["B P CODE", "S P CODE", "SHIP PLANT", "MATERIAL", "DESCRIPTION"]
        if any(col not in idx for col in required):
            return {}

        for row in ws.iter_rows(min_row=2, values_only=True):
            bp = normalise_party_code(row[idx["B P CODE"]] if idx["B P CODE"] < len(row) else "")
            sp = normalise_party_code(row[idx["S P CODE"]] if idx["S P CODE"] < len(row) else "")
            plant = normalise_plant_code(row[idx["SHIP PLANT"]] if idx["SHIP PLANT"] < len(row) else "")
            material = str(row[idx["MATERIAL"]] if idx["MATERIAL"] < len(row) else "").strip().upper()
            desc = str(row[idx["DESCRIPTION"]] if idx["DESCRIPTION"] < len(row) else "").strip().upper()
            if bp != target_bp or sp != target_sp or plant != target_plant or material != target_material or desc != target_desc:
                continue

            def col(name: str) -> str:
                if name not in idx or idx[name] >= len(row):
                    return ""
                value = row[idx[name]]
                return "" if value is None else str(value).strip()

            return {
                "customer_order_category": col("CUST ORDER"),
                "eq_specif_grp": col("Eq.Specif.Grp") or col("EqSpecifGrp"),
                "eq_specifi": col("Eq.Specifi.") or col("EqSpecifi"),
                "eq_sub_grade": col("Eq.Sub_Grade") or col("EqSub_Grade"),
                "end_appn": col("END_APPN"),
                "rh_req": col("RH REQ"),
                "plant_code": plant,
                "width": col("WIDTH"),
                "thickness": col("THICKNESS"),
                "length": col("LENGTH"),
                "edge_con": col("EDGE_CON"),
                "thick_tol_type": col("THICK_TOL_TYPE"),
                "oil_req": col("OIL_REQ"),
                "s_brand": col("BRAND"),
                "spangle_type": col("S_SPANGLE_TYPE"),
                "zinc_coating_min": col("ZINC COATING") or col("ZINC_COAT"),
                "description": str(payload.get("description") or "").strip(),
                "material": str(payload.get("material") or "").strip(),
            }
    except Exception:
        return {}
    return {}


def _call_lookup(payload: dict[str, Any]) -> dict[str, Any]:
    division = str(payload.get("division") or "").strip().upper()
    url = _lookup_url_for_division(division)
    if not url:
        raise RuntimeError(f"{_lookup_env_for_division(division)} is not configured.")

    try:
        response = requests.post(url, json=payload, timeout=90)
        response.raise_for_status()
    except requests.RequestException:
        return {"status": "success", "materials": [], "skus": [], "rows": []}

    if not (response.text or "").strip():
        return {"status": "success", "materials": [], "skus": [], "rows": []}

    try:
        data = response.json()
    except ValueError:
        return {"status": "success", "materials": [], "skus": [], "rows": []}

    if str(data.get("status", "success")).lower() not in {"success", "ok"}:
        raise RuntimeError(data.get("message") or data.get("error") or f"{division or 'SKU'} master lookup failed.")
    for key in ("materials", "skus", "rows"):
        data[key] = _normalise_list_value(data.get(key))
    data["rows"] = _normalise_row_items(data.get("rows"))
    return data


def _normalise_list_value(value: Any) -> list[Any]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return []
        try:
            parsed = json.loads(text)
            if isinstance(parsed, list):
                return parsed
            return [parsed]
        except ValueError:
            return [value]
    return [value]


def _normalise_row_items(rows: Any) -> list[dict[str, Any]]:
    normalised: list[dict[str, Any]] = []
    for item in rows or []:
        if isinstance(item, dict):
            normalised.append(item)
            continue
        parsed = _try_parse_object(item)
        if isinstance(parsed, dict):
            normalised.append(parsed)
    return normalised


def _try_parse_object(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        return value
    text = str(value).strip()
    if not text:
        return None
    # Handle both JSON and single-quoted pseudo-json payloads.
    for parser in (json.loads, ast.literal_eval):
        try:
            return parser(text)
        except Exception:
            continue
    return None


def _lookup_url_for_division(division: str) -> str:
    return os.getenv(_lookup_env_for_division(division), "").strip()


def _lookup_env_for_division(division: str) -> str:
    if division == "CRCA":
        return "CRCA_MASTER_LOOKUP_URL"
    if division == "GI":
        return "GI_MASTER_LOOKUP_URL"
    if division == "GL":
        return "GL_MASTER_LOOKUP_URL"
    return "HRC_MASTER_LOOKUP_URL"
